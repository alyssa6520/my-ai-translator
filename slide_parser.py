import os
import re
from typing import List, Dict, Optional
from dataclasses import dataclass, asdict
from pypdf import PdfReader
from pptx import Presentation


@dataclass
class SlideContent:
    page_num: int
    title: str
    content: str
    full_text: str
    images: List[str]

    def to_dict(self):
        return asdict(self)


class SlideParser:
    def __init__(self):
        pass

    @staticmethod
    def _clean_text(text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    @staticmethod
    def _extract_title(text: str, fallback: str) -> str:
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        if lines:
            first = lines[0]
            if len(first) < 100:
                return first
        return fallback

    def parse_pdf(self, pdf_path: str) -> List[SlideContent]:
        slides = []
        reader = PdfReader(pdf_path)
        total_pages = len(reader.pages)

        for i, page in enumerate(reader.pages):
            page_num = i + 1
            text = page.extract_text() or ""
            title = self._extract_title(text, f"第 {page_num} 页")
            cleaned = self._clean_text(text)
            slides.append(SlideContent(
                page_num=page_num,
                title=title,
                content=cleaned,
                full_text=text,
                images=[]
            ))
        return slides

    def parse_pptx(self, pptx_path: str) -> List[SlideContent]:
        slides = []
        prs = Presentation(pptx_path)

        for i, slide in enumerate(prs.slides):
            page_num = i + 1
            texts = []
            title = f"第 {page_num} 页"

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for j, para in enumerate(shape.text_frame.paragraphs):
                        para_text = "".join(run.text for run in para.runs)
                        if para_text.strip():
                            if j == 0 and len(para_text.strip()) < 100 and shape == slide.shapes[0]:
                                title = para_text.strip()
                            texts.append(para_text)

            full_text = "\n".join(texts)
            cleaned = self._clean_text(full_text)
            slides.append(SlideContent(
                page_num=page_num,
                title=title,
                content=cleaned,
                full_text=full_text,
                images=[]
            ))
        return slides

    def parse_file(self, file_path: str) -> List[SlideContent]:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self.parse_pdf(file_path)
        elif ext in [".pptx", ".ppt"]:
            return self.parse_pptx(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    def get_all_texts(self, slides: List[SlideContent]) -> List[str]:
        return [f"{s.title} {s.content}" for s in slides]
